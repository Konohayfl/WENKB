"""Clear historical failure executions while preserving cases and screenshots."""

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
WORKBOOK = (
    ROOT
    / "outputs"
    / "wenkb-test-book-in-progress"
    / "第14小组_基于大模型的个人知识库系统_测试成绩书.xlsx"
)
FAILED_CASES = {
    "TC-WKB-008",
    "TC-WKB-010",
    "TC-WKB-015",
    "TC-WKB-038",
    "TC-WKB-041",
    "TC-WKB-042",
    "TC-WKB-043",
    "TC-WKB-044",
}


def main() -> None:
    workbook = load_workbook(WORKBOOK)
    cases = workbook["测试项目"]
    run = workbook["运行"]
    image_count = len(run._images)
    seen = set()

    for row in range(1, cases.max_row + 1):
        case_id = cases.cell(row, 1).value
        if case_id not in FAILED_CASES:
            continue
        if cases.cell(row, 9).value not in {"失败", "未执行（待复测）"}:
            raise ValueError(f"Unexpected status for {case_id}")
        seen.add(case_id)
        cases.cell(row, 8).value = None
        cases.cell(row, 9).value = "未执行（待复测）"
        cases.cell(row, 10).value = None
        cases.cell(row, 11).value = None
        cases.cell(row, 12).value = None
        cases.cell(row, 13).value = None

    if seen != FAILED_CASES:
        raise ValueError(f"Missing expected failed cases: {FAILED_CASES - seen}")

    cover = workbook["文件封面"]
    cover["B14"] = "失败执行记录已转存，8 条待复测；阻塞项待手动执行"
    cover["B18"] = "18 通过，0 失败，22 阻塞，8 待复测；截图已粘贴"
    run["A36"] = "项目实际运行情况（截图已粘贴；详见用例执行状态）"
    run["A37"] = (
        "原执行中的 8 条失败详情已转存到 docs/test_execution_followup.md，"
        "工作簿对应执行字段清空并待复测；22 条阻塞保留待手动执行。"
    )
    run["B14"] = None
    run["B19"] = "WenKB 浏览器工作区与主导航"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(WORKBOOK)

    verified = load_workbook(WORKBOOK, data_only=False)
    verified_cases = verified["测试项目"]
    statuses = [
        verified_cases.cell(row, 9).value
        for row in range(1, verified_cases.max_row + 1)
        if str(verified_cases.cell(row, 1).value or "").startswith("TC-")
    ]
    expected = {"通过": 18, "阻塞": 22, "未执行（待复测）": 8}
    actual = {status: statuses.count(status) for status in expected}
    if actual != expected:
        raise AssertionError(f"Unexpected status totals: {actual}")
    if len(verified["运行"]._images) != image_count:
        raise AssertionError("Embedded runtime screenshots changed")
    print(f"Updated {WORKBOOK}: {actual}; preserved {image_count} screenshots")


if __name__ == "__main__":
    main()
